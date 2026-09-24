"""Ustvari Excel program za beleženje kilometrine službenih poti (KILOMETRINA.xlsx).

Izgled sledi vzorcu "PODPISNI LIST - MED ENOTAMI": naslov, ime, mesec,
tabela DOGODEK / RELACIJA / KILOMETRI / VREDNOST, vrstica SKUPAJ in opombe,
ob strani pa tabela kilometrov med enotami in kilometrina.
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

FONT = "Calibri"
OUT = "KILOMETRINA.xlsx"

MESECI = ["Januar", "Februar", "Marec", "April", "Maj", "Junij", "Julij",
          "Avgust", "September", "Oktober", "November", "December"]

# Barve iz vzorca (svetle različice tem barv)
ORANZNA = PatternFill("solid", fgColor="FBE4D5")
SIVA = PatternFill("solid", fgColor="EDEDED")
MODRA = PatternFill("solid", fgColor="DDEBF7")
BEZ = PatternFill("solid", fgColor="E7E6E6")
VNOS = PatternFill("solid", fgColor="FFFF99")  # celice za vnos

thin = Side(style="thin")
med = Side(style="medium")
TANKA = Border(left=thin, right=thin, top=thin, bottom=thin)
SREDNJA = Border(left=med, right=med, top=med, bottom=med)

EUR = '#,##0.00 "€"'
KM = '#,##0.0 "km"'
DATUM = "d.m.yyyy"

# Relacije: (naziv, km v eno smer) - vrednosti iz vzorca
RELACIJE = [
    ("Šmartno pri Cerkljah - Križe", 17.45),
    ("Šmartno pri Cerkljah - Palček", 14.8),
    ("Šmartno pri Cerkljah - Lom", 18),
    ("Šmartno pri Cerkljah - Deteljica", 16.5),
    ("Križe - Muzej", 3.5),
    ("Križe - Kovor", 1.6),
    ("Palček - Občina", 1.1),
    ("Palček - Deteljica", 1.2),
    ("Palček - Kovor", 2.8),
    ("Palček - Križe", 2.7),
    ("Palček - OŠ Tržič", 1.5),
    ("Palček - Lom", 3.3),
    ("Palček - Knjižnica", 1.1),
    ("Palček - Muzej", 1.5),
    ("Deteljica - Kovor", 1.9),
    ("Deteljica - Muzej", 2.6),
    ("Deteljica - Križe", 2.7),
]
REL_VRSTIC = 100  # prostor za nove relacije
SMERI = ["tja in nazaj", "ena smer"]
VRSTIC_MESEC = 31
PRVA = 10  # prva vrstica podatkov na mesečnem listu


def celica(ws, ref, value=None, *, bold=False, size=11, fill=None, border=None,
           h=None, v="center", wrap=False, fmt=None, color=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name=FONT, bold=bold, size=size, color=color)
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if fmt:
        c.number_format = fmt
    return c


wb = Workbook()

# ---------------------------------------------------------------- NASTAVITVE
nst = wb.active
nst.title = "Nastavitve"
nst.sheet_view.showGridLines = False
for col, w in {"A": 2, "B": 34, "C": 34, "D": 3, "E": 16, "F": 14}.items():
    nst.column_dimensions[col].width = w

celica(nst, "B1", "NASTAVITVE", bold=True, size=14)
celica(nst, "B2", "Rumene celice izpolnite. Podatki se samodejno prenesejo na vse mesečne liste.",
       size=10, color="595959")

podatki = [
    ("Ime in priimek", "Ime Priimek"),
    ("Delovno mesto / enota", "Enota"),
    ("Enote (izpis pod mesecem)", "PALČEK, KRIŽE, DETELJICA, LOM, KOVOR"),
    ("Leto", 2026),
]
for i, (label, val) in enumerate(podatki, start=4):
    celica(nst, f"B{i}", label, border=TANKA, fill=BEZ)
    celica(nst, f"C{i}", val, border=TANKA, fill=VNOS, h="left")
for name, ref in [("Ime", "$C$4"), ("DelovnoMesto", "$C$5"), ("Enote", "$C$6"), ("Leto", "$C$7")]:
    wb.defined_names[name] = DefinedName(name, attr_text=f"Nastavitve!{ref}")

# Kilometrina (kot v vzorcu: velja od datuma)
celica(nst, "E3", "KILOMETRINA", bold=True, size=12, fill=SIVA, border=TANKA, h="center")
nst.merge_cells("E3:F3")
celica(nst, "E4", "velja od", fill=SIVA, border=TANKA, h="center")
celica(nst, "F4", "€ / km", fill=SIVA, border=TANKA, h="center")
celica(nst, "E5", "=DATE(Leto,1,1)", fill=VNOS, border=TANKA, h="center", fmt=DATUM)
celica(nst, "F5", 0.43, fill=VNOS, border=TANKA, h="center", fmt='#,##0.00"€"')
nst["F5"].comment = Comment(
    "Predpostavka: 0,43 €/km. Preverite veljavno kilometrino pri delodajalcu "
    "in vrednost po potrebi popravite.", "Kilometrina")
for r in range(6, 11):
    celica(nst, f"E{r}", fill=VNOS, border=TANKA, h="center", fmt=DATUM)
    celica(nst, f"F{r}", fill=VNOS, border=TANKA, h="center", fmt='#,##0.00"€"')
celica(nst, "E11", "Ob spremembi vpišite nov datum in znesek v naslednjo vrstico (datumi naraščajoče).",
       size=9, color="595959", wrap=True, v="top")
nst.merge_cells("E11:F13")
wb.defined_names["KmDatumi"] = DefinedName("KmDatumi", attr_text="Nastavitve!$E$5:$E$10")
wb.defined_names["KmCene"] = DefinedName("KmCene", attr_text="Nastavitve!$F$5:$F$10")

# Dela prosti dnevi v Sloveniji - za izračun delovnih dni
# Velikonočna nedelja: znana formula, veljavna za leta 1900-2203
VELIKA_NOC = "(FLOOR(DATE(Leto,5,DAY(MINUTE(Leto/38)/2+56)),7)-34)"
PRAZNIKI = [
    ("Novo leto", "=DATE(Leto,1,1)"),
    ("Novo leto", "=DATE(Leto,1,2)"),
    ("Prešernov dan", "=DATE(Leto,2,8)"),
    ("Velikonočni ponedeljek", f"={VELIKA_NOC}+1"),
    ("Dan upora proti okupatorju", "=DATE(Leto,4,27)"),
    ("Praznik dela", "=DATE(Leto,5,1)"),
    ("Praznik dela", "=DATE(Leto,5,2)"),
    ("Dan državnosti", "=DATE(Leto,6,25)"),
    ("Marijino vnebovzetje", "=DATE(Leto,8,15)"),
    ("Dan reformacije", "=DATE(Leto,10,31)"),
    ("Dan spomina na mrtve", "=DATE(Leto,11,1)"),
    ("Božič", "=DATE(Leto,12,25)"),
    ("Dan samostojnosti in enotnosti", "=DATE(Leto,12,26)"),
]
PRAZ_VRSTIC = len(PRAZNIKI) + 7  # prostor za dodatne dela proste dneve
celica(nst, "B15", "DELA PROSTI DNEVI (PRAZNIKI)", bold=True, size=12, fill=SIVA,
       border=TANKA, h="center")
nst.merge_cells("B15:C15")
celica(nst, "B16", "praznik", fill=SIVA, border=TANKA, h="center")
celica(nst, "C16", "datum", fill=SIVA, border=TANKA, h="center")
for i in range(PRAZ_VRSTIC):
    r = 17 + i
    naziv, formula = PRAZNIKI[i] if i < len(PRAZNIKI) else (None, None)
    celica(nst, f"B{r}", naziv, border=TANKA, fill=None if naziv else VNOS)
    celica(nst, f"C{r}", formula, border=TANKA, h="center", fmt="ddd d.m.yyyy",
           fill=None if formula else VNOS)
zadnji_praz = 16 + PRAZ_VRSTIC
celica(nst, f"B{zadnji_praz + 1}", "Prazniki se samodejno izračunajo za izbrano leto "
       "(velika noč po formuli). V rumene vrstice lahko dodate druge dela proste dneve.",
       size=9, color="595959", wrap=True, v="top")
nst.merge_cells(f"B{zadnji_praz + 1}:C{zadnji_praz + 2}")
wb.defined_names["Prazniki"] = DefinedName("Prazniki", attr_text=f"Nastavitve!$C$17:$C${zadnji_praz}")

# ------------------------------------------------------------------ RELACIJE
rel = wb.create_sheet("Relacije")
rel.sheet_view.showGridLines = False
for col, w in {"A": 2, "B": 42, "C": 14, "D": 14}.items():
    rel.column_dimensions[col].width = w
celica(rel, "B1", "KILOMETRI MED ENOTAMI / RELACIJE", bold=True, size=14)
celica(rel, "B2", "Vpišite relacijo in kilometre v eno smer (rumeno). Obe smeri se izračunata samodejno.",
       size=10, color="595959")
for col, t in zip("BCD", ["RELACIJA", "KILOMETRI ENA SMER", "KILOMETRI OBE SMERI"]):
    celica(rel, f"{col}4", t, bold=True, fill=MODRA, border=TANKA, h="center", wrap=True)
rel.row_dimensions[4].height = 30
for i in range(REL_VRSTIC):
    r = 5 + i
    naziv, km = RELACIJE[i] if i < len(RELACIJE) else (None, None)
    celica(rel, f"B{r}", naziv, fill=VNOS, border=TANKA)
    celica(rel, f"C{r}", km, fill=VNOS, border=TANKA, h="center", fmt="0.0#")
    celica(rel, f"D{r}", f'=IF(C{r}="","",C{r}*2)', border=TANKA, h="center", fmt="0.0#")
last = 4 + REL_VRSTIC
rel.freeze_panes = "B5"
wb.defined_names["RelNazivi"] = DefinedName("RelNazivi", attr_text=f"Relacije!$B$5:$B${last}")
wb.defined_names["RelKm"] = DefinedName("RelKm", attr_text=f"Relacije!$C$5:$C${last}")

# ------------------------------------------------------------- MESEČNI LISTI
ZADNJA = PRVA + VRSTIC_MESEC - 1
SKUPAJ = ZADNJA + 1

for m, mesec in enumerate(MESECI, start=1):
    ws = wb.create_sheet(mesec)
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 13, "C": 40, "D": 34, "E": 14, "F": 13, "G": 14,
                   "H": 3, "I": 12, "J": 12}.items():
        ws.column_dimensions[col].width = w

    celica(ws, "B1", "PODPISNI LIST - KILOMETRINA SLUŽBENIH POTI", bold=True, size=14, h="left")
    ws.merge_cells("B1:E1")
    celica(ws, "B3", "=Ime", bold=True, size=12, h="left")
    ws.merge_cells("B3:D3")
    celica(ws, "B4", "=DelovnoMesto", size=12, h="left")
    ws.merge_cells("B4:D4")
    celica(ws, "B5", f"=DATE(Leto,{m},1)", bold=True, size=12, h="left", fmt="mmmm yyyy")
    ws.merge_cells("B5:C5")
    celica(ws, "B6", "=Enote", size=12, h="left")
    ws.merge_cells("B6:D6")

    # Kilometrina (desno zgoraj, kot v vzorcu)
    celica(ws, "I2", "kilometrina", border=TANKA, h="center")
    ws.merge_cells("I2:J2")
    ws["J2"].border = TANKA
    celica(ws, "I3", "€ / km", fill=SIVA, border=TANKA, h="center")
    celica(ws, "J3", f"=INDEX(KmCene,MATCH(DATE(Leto,{m},1),KmDatumi,1))",
           fill=SIVA, border=TANKA, h="center", fmt='#,##0.00"€"')
    ws["J3"].comment = Comment("Kilometrina, ki velja na 1. dan meseca (list Nastavitve).",
                               "Kilometrina")
    celica(ws, "I4", "delovni dnevi", fill=MODRA, border=TANKA, h="center")
    celica(ws, "J4", f"=NETWORKDAYS(DATE(Leto,{m},1),EOMONTH(DATE(Leto,{m},1),0),Prazniki)",
           fill=MODRA, border=TANKA, h="center", bold=True)
    ws["J4"].comment = Comment("Število delovnih dni v mesecu (pon-pet, brez praznikov "
                               "s lista Nastavitve).", "Delovni dnevi")

    glave = ["DATUM", "DOGODEK / NAMEN POTI", "RELACIJA*", "POT", "KILOMETRI", "VREDNOST**"]
    for col, t in zip("BCDEFG", glave):
        celica(ws, f"{col}{PRVA - 1}", t, bold=True, size=12, border=SREDNJA,
               h="center", wrap=True, fill=ORANZNA)
    ws.row_dimensions[PRVA - 1].height = 33

    for r in range(PRVA, ZADNJA + 1):
        ws.row_dimensions[r].height = 18
        # n-ti delovni dan meseca (brez vikendov in praznikov); po zadnjem ostane prazno
        n = r - PRVA + 1
        dat = (f'=IF(WORKDAY(DATE(Leto,{m},1)-1,{n},Prazniki)>EOMONTH(DATE(Leto,{m},1),0),'
               f'"",WORKDAY(DATE(Leto,{m},1)-1,{n},Prazniki))')
        celica(ws, f"B{r}", dat, border=SREDNJA, h="center", fmt=DATUM)
        celica(ws, f"C{r}", border=SREDNJA, wrap=True)
        celica(ws, f"D{r}", border=SREDNJA)
        celica(ws, f"E{r}", border=SREDNJA, h="center")
        celica(ws, f"F{r}",
               f'=IF(D{r}="","",IFERROR(INDEX(RelKm,MATCH(D{r},RelNazivi,0))'
               f'*IF(E{r}="ena smer",1,2),""))',
               border=SREDNJA, h="center", fmt=KM)
        celica(ws, f"G{r}", f'=IF(F{r}="","",ROUND(F{r}*$J$3,2))',
               border=SREDNJA, h="center", fmt=EUR)

    celica(ws, f"B{SKUPAJ}", "SKUPAJ", bold=True, size=16, border=SREDNJA, fill=ORANZNA)
    ws.merge_cells(f"B{SKUPAJ}:E{SKUPAJ}")
    for col in "CDE":
        ws[f"{col}{SKUPAJ}"].border = SREDNJA
    celica(ws, f"F{SKUPAJ}", f"=SUM(F{PRVA}:F{ZADNJA})", bold=True, size=12,
           border=SREDNJA, h="center", fmt=KM, fill=ORANZNA)
    celica(ws, f"G{SKUPAJ}", f"=SUM(G{PRVA}:G{ZADNJA})", bold=True, size=12,
           border=SREDNJA, h="center", fmt=EUR, fill=ORANZNA)
    ws.row_dimensions[SKUPAJ].height = 24

    celica(ws, f"B{SKUPAJ + 1}", "*relacija se vpisuje od enote, kjer opravljate delo, do enote, "
           "kjer se odvija dogodek (izbira s seznama na listu Relacije)", size=10, h="left")
    celica(ws, f"B{SKUPAJ + 2}", "**do povračila potnih stroškov so upravičeni zaposleni, ki v "
           "tekočem mesecu niso prejeli povračila potnih stroškov (mesečna karta) oziroma prejemajo "
           "povračilo potnih stroškov na drugi relaciji", size=10, h="left")
    ws.merge_cells(f"B{SKUPAJ + 2}:G{SKUPAJ + 3}")
    ws[f"B{SKUPAJ + 2}"].alignment = Alignment(wrap_text=True, vertical="top")

    pod = SKUPAJ + 6
    celica(ws, f"B{pod}", "Datum:", size=11)
    celica(ws, f"C{pod}", border=Border(bottom=thin))
    celica(ws, f"D{pod}", "Podpis zaposlenega:", size=11, h="right")
    ws.merge_cells(f"E{pod}:G{pod}")
    for col in "EFG":
        ws[f"{col}{pod}"].border = Border(bottom=thin)
    celica(ws, f"D{pod + 2}", "Odobril:", size=11, h="right")
    ws.merge_cells(f"E{pod + 2}:G{pod + 2}")
    for col in "EFG":
        ws[f"{col}{pod + 2}"].border = Border(bottom=thin)

    # Spustni seznami in preverjanje datuma
    dv_rel = DataValidation(type="list", formula1="=RelNazivi", allow_blank=True,
                            showErrorMessage=True, errorTitle="Relacija",
                            error="Izberite relacijo s seznama (novo dodajte na list Relacije).")
    dv_smer = DataValidation(type="list", formula1='"' + ",".join(SMERI) + '"', allow_blank=True)
    dv_dat = DataValidation(type="date", operator="between",
                            formula1=f"DATE(Leto,{m},1)", formula2=f"EOMONTH(DATE(Leto,{m},1),0)",
                            allow_blank=True, showErrorMessage=True, errorTitle="Datum",
                            error=f"Datum mora biti v mesecu {mesec.lower()}.")
    for dv in (dv_rel, dv_smer, dv_dat):
        ws.add_data_validation(dv)
    dv_dat.add(f"B{PRVA}:B{ZADNJA}")
    dv_rel.add(f"D{PRVA}:D{ZADNJA}")
    dv_smer.add(f"E{PRVA}:E{ZADNJA}")

    ws.freeze_panes = f"B{PRVA}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:J{pod + 2}"

# ----------------------------------------------------------- LETNI PREGLED
lp = wb.create_sheet("Letni pregled", 2)
lp.sheet_view.showGridLines = False
for col, w in {"A": 2, "B": 18, "C": 14, "D": 16, "E": 16, "F": 16}.items():
    lp.column_dimensions[col].width = w
celica(lp, "B1", "LETNI PREGLED KILOMETRINE", bold=True, size=14)
celica(lp, "B2", '="Leto "&Leto&" - "&Ime', size=12)
for col, t in zip("BCDEF", ["MESEC", "ŠT. POTI", "KILOMETRI", "VREDNOST", "DELOVNI DNEVI"]):
    celica(lp, f"{col}4", t, bold=True, size=12, border=SREDNJA, fill=ORANZNA, h="center")
for i, mesec in enumerate(MESECI):
    r = 5 + i
    celica(lp, f"B{r}", mesec, border=SREDNJA)
    celica(lp, f"C{r}", f"=COUNT({mesec}!F{PRVA}:F{ZADNJA})", border=SREDNJA, h="center",
           color="008000")
    celica(lp, f"D{r}", f"={mesec}!F{SKUPAJ}", border=SREDNJA, h="center", fmt=KM, color="008000")
    celica(lp, f"E{r}", f"={mesec}!G{SKUPAJ}", border=SREDNJA, h="center", fmt=EUR, color="008000")
    celica(lp, f"F{r}", f"={mesec}!J4", border=SREDNJA, h="center", color="008000")
celica(lp, "B17", "SKUPAJ", bold=True, size=14, border=SREDNJA, fill=ORANZNA)
for col, fmt in zip("CDEF", ["0", KM, EUR, "0"]):
    celica(lp, f"{col}17", f"=SUM({col}5:{col}16)", bold=True, size=12, border=SREDNJA,
           fill=ORANZNA, h="center", fmt=fmt)

# ---------------------------------------------------------------- NAVODILA
nav = wb.create_sheet("Navodila", 0)
nav.sheet_view.showGridLines = False
nav.column_dimensions["A"].width = 2
nav.column_dimensions["B"].width = 110
vrstice = [
    ("NAVODILA ZA UPORABO", True, 14),
    ("", False, 11),
    ("1. List Nastavitve: vpišite ime in priimek, delovno mesto, enote, leto in veljavno kilometrino (€/km).", False, 11),
    ("2. List Relacije: vpišite relacije in kilometre v ENO smer. Seznam lahko poljubno dopolnjujete (do 100 relacij).", False, 11),
    ("3. Mesečni listi (Januar ... December): stolpec DATUM je že izpolnjen z delovnimi dnevi meseca", False, 11),
    ("     (brez sobot, nedelj in praznikov s lista Nastavitve). Za dan s službeno potjo vpišite", False, 11),
    ("     DOGODEK / NAMEN POTI  →  DOGODEK / NAMEN POTI  →  izberite RELACIJO s seznama  →  izberite POT (tja in nazaj / ena smer).", False, 11),
    ("     KILOMETRI in VREDNOST se izračunata samodejno (prazna POT pomeni tja in nazaj).", False, 11),
    ("     Pot na dela prost dan ali drugo pot v istem dnevu vpišite v prazne vrstice pod zadnjim delovnim dnem.", False, 11),
    ("4. List Letni pregled samodejno povzame število poti, kilometre in znesek po mesecih.", False, 11),
    ("5. Mesečni list je pripravljen za tisk na A4 ležeče (podpisni list).", False, 11),
    ("", False, 11),
    ("Rumene celice = vnos podatkov. Ostale celice vsebujejo formule - ne prepisujte jih.", True, 11),
    ("", False, 11),
    ("PRIMER VNOSA:", True, 11),
    ("     1.9.2026  |  Popoldanska delavnica  |  Šmartno pri Cerkljah - Križe  |  tja in nazaj  →  34,9 km  |  15,01 €  (pri 0,43 €/km)", False, 11),
]
for i, (t, b, s) in enumerate(vrstice, start=1):
    celica(nav, f"B{i}", t, bold=b, size=s, h="left")
for c in nav["B"]:
    if str(c.value).startswith("Rumene celice"):
        c.fill = VNOS

wb.active = wb.sheetnames.index("Januar")
wb.save(OUT)
print("Shranjeno:", OUT)
